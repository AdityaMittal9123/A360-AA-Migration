"""A360 -> Excel standalone converter.

One self-contained script: point it at a Control Room export (.zip), a folder of
bot JSONs, or a single .json, and it parses the estate and writes a 4-sheet .xlsx
(Summary, Call Tree, Task bots, Task bot metrics) straight to your laptop.

Run:
    pip install openpyxl
    python parser_to_excel.py

Then follow the prompts. Press Enter at the path prompt to use the bundled sample.
"""
import io
import json
import os
import re
import sys
import zipfile
from collections import defaultdict
from datetime import datetime

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill

# ---------------------------------------------------------------------------
# Complexity model (edit these weights — they are deliberately visible)
# ---------------------------------------------------------------------------
LOC_BANDS = [            # (max_active_actions, score)
    (100, 1), (200, 2), (400, 3), (700, 4), (10**9, 5)
]
PACKAGE_WEIGHT = {       # per action, by package family
    "variable": 0.2, "string": 0.2, "number": 0.2, "datetime": 0.2, "boolean": 0.2,
    "if": 0.4, "loop": 0.6, "errorhandler": 0.5, "step": 0.1, "comment": 0.0,
    "taskbot": 0.5, "dictionary": 0.4, "list": 0.4, "datatable": 0.6,
    "email": 1.0, "rest": 1.2, "database": 1.2, "http": 1.2, "soap": 1.2,
    "excel": 1.0, "csv": 0.8, "file": 0.6, "folder": 0.6, "pdf": 1.0,
    "recorder": 2.5, "objectcloning": 2.5, "sap": 2.5, "browser": 2.0,
    "terminal": 2.5, "image": 3.0, "ocr": 2.0, "citrix": 3.0, "window": 1.5,
    "iqbot": 2.0, "documentautomation": 2.0,
    "dll": 3.0, "python": 3.0, "vbscript": 3.0, "powershell": 2.5, "javascript": 2.5,
    "wlm": 1.5, "workload": 1.5, "credential": 0.8, "system": 0.8, "clipboard": 0.6,
}
DEFAULT_WEIGHT = 1.0
COMPLEXITY_THRESHOLDS = [("Low", 5.0), ("Medium", 10.0), ("High", 10**9)]

# ---------------------------------------------------------------------------
# Loading
# ---------------------------------------------------------------------------
def _load_bot_files(path):
    """Yield (relative_name, json_obj) for every bot definition found."""
    if os.path.isdir(path):
        for root, _, files in os.walk(path):
            for f in files:
                if f.lower().endswith((".json", ".abot")):
                    with open(os.path.join(root, f), "r", encoding="utf-8", errors="ignore") as fh:
                        obj = _try_json(fh.read())
                        if obj is not None:
                            yield os.path.relpath(os.path.join(root, f), path), obj
    elif path.lower().endswith(".zip"):
        with zipfile.ZipFile(path) as z:
            for name in z.namelist():
                if name.lower().endswith((".json", ".abot")) and not name.endswith("/"):
                    obj = _try_json(z.read(name).decode("utf-8", errors="ignore"))
                    if obj is not None:
                        yield name, obj
    else:
        with open(path, "r", encoding="utf-8", errors="ignore") as fh:
            obj = _try_json(fh.read())
            if obj is not None:
                yield os.path.basename(path), obj


def _try_json(text):
    try:
        return json.loads(text)
    except Exception:
        return None


def _is_bot(obj):
    return isinstance(obj, dict) and isinstance(obj.get("nodes"), list)


# ---------------------------------------------------------------------------
# Walking a bot
# ---------------------------------------------------------------------------
def _walk(nodes, depth, out, disabled_parent=False):
    for n in nodes or []:
        if not isinstance(n, dict):
            continue
        pkg = (n.get("packageName") or n.get("package") or "").strip()
        cmd = (n.get("commandName") or n.get("command") or "").strip()
        disabled = bool(n.get("disabled", False)) or disabled_parent
        attrs = {}
        for a in n.get("attributes", []) or []:
            if isinstance(a, dict) and "name" in a:
                v = a.get("value")
                if isinstance(v, dict):          # A360 wraps values: {"type":"STRING","string":"..."}
                    v = v.get("string") or v.get("number") or v.get("expression") or json.dumps(v)
                attrs[a["name"]] = v
        out.append({
            "uid": n.get("uid"), "package": pkg, "command": cmd,
            "disabled": disabled, "depth": depth, "attributes": attrs,
        })
        # children live under different keys depending on the command
        for key in ("children", "branches", "nodes"):
            if isinstance(n.get(key), list):
                for child in n[key]:
                    if isinstance(child, dict) and isinstance(child.get("nodes"), list):
                        _walk(child["nodes"], depth + 1, out, disabled)   # If/Else branches
                    else:
                        _walk([child], depth + 1, out, disabled)


def _sub_task_path(action):
    a = action["attributes"]
    for k in ("botPath", "path", "fileId", "botName", "taskbot"):
        if k in a and a[k]:
            return str(a[k])
    return None


def parse_bot(name, obj):
    actions = []
    _walk(obj.get("nodes", []), 0, actions)
    active = [a for a in actions if not a["disabled"]]
    disabled = [a for a in actions if a["disabled"]]

    packages = defaultdict(int)
    for a in active:
        packages[a["package"] or "unknown"] += 1

    calls = []
    for a in active:
        if a["package"].lower() == "taskbot" and a["command"].lower() in ("run", "runtask", "run_task"):
            p = _sub_task_path(a)
            if p:
                calls.append(p)

    variables = []
    for v in obj.get("variables", []) or []:
        if isinstance(v, dict):
            variables.append({"name": v.get("name"), "type": v.get("type"),
                              "input": bool(v.get("input")), "output": bool(v.get("output"))})

    vault_refs = sorted({m for a in actions for val in a["attributes"].values()
                         for m in re.findall(r"\$Vault\.[A-Za-z0-9_\.]+\$", str(val))})
    global_refs = sorted({m for a in actions for val in a["attributes"].values()
                          for m in re.findall(r"\$Global\.[A-Za-z0-9_]+\$", str(val))})
    max_depth = max((a["depth"] for a in actions), default=0)

    score = _score(len(active), packages, max_depth)
    return {
        "name": _bot_name(name, obj),
        "file": name,
        "active_actions": len(active),
        "disabled_actions": len(disabled),
        "total_actions": len(actions),
        "max_nesting_depth": max_depth,
        "variables": variables,
        "packages": dict(sorted(packages.items(), key=lambda kv: -kv[1])),
        "calls": calls,
        "called_by": [],
        "vault_references": vault_refs,
        "global_value_references": global_refs,
        "complexity_score": round(score, 1),
        "complexity": _band(score),
        "actions": actions,      # full normalised action list
    }


def _bot_name(fname, obj):
    for k in ("name", "botName", "displayName"):
        if obj.get(k):
            return str(obj[k])
    return os.path.splitext(os.path.basename(fname))[0]


def _score(active_count, packages, depth):
    loc = next(s for mx, s in LOC_BANDS if active_count <= mx)
    weighted = 0.0
    for pkg, cnt in packages.items():
        key = pkg.lower().replace(" ", "")
        w = next((v for k, v in PACKAGE_WEIGHT.items() if k in key), DEFAULT_WEIGHT)
        weighted += w * cnt
    # normalise the weighted action score to a 0–10 band, add LOC and nesting
    return loc + min(weighted / 8.0, 10.0) + min(depth * 0.5, 3.0)


def _band(score):
    return next(label for label, mx in COMPLEXITY_THRESHOLDS if score <= mx)


# ---------------------------------------------------------------------------
# Estate
# ---------------------------------------------------------------------------
def parse_export(path):
    bots = []
    for name, obj in _load_bot_files(path):
        if _is_bot(obj):
            bots.append(parse_bot(name, obj))
    # resolve calls → called_by (match on file path tail or bot name)
    by_key = {}
    for b in bots:
        by_key[b["name"].lower()] = b
        by_key[os.path.splitext(os.path.basename(b["file"]))[0].lower()] = b
    for b in bots:
        resolved = []
        for c in b["calls"]:
            tail = os.path.splitext(os.path.basename(c.replace("\\", "/")))[0].lower()
            target = by_key.get(tail) or by_key.get(c.lower())
            resolved.append(target["name"] if target else c)
            if target:
                target["called_by"].append(b["name"])
        b["calls"] = resolved
    roots = [b["name"] for b in bots if not b["called_by"]]
    return {
        "source": os.path.basename(path),
        "task_bots": len(bots),
        "active_actions": sum(b["active_actions"] for b in bots),
        "disabled_actions": sum(b["disabled_actions"] for b in bots),
        "roots": roots,
        "bots": bots,
    }


# ---------------------------------------------------------------------------
# Excel building (4 sheets mirroring the Parse screen)
# ---------------------------------------------------------------------------
HEADER_FILL = PatternFill("solid", fgColor="1F3B57")
HEADER_FONT = Font(bold=True, color="FFFFFF")
SECTION_FONT = Font(bold=True)
LABEL_FONT = Font(bold=True, color="555555")


def _header(ws, cols, widths):
    ws.append(cols)
    for i, (c, w) in enumerate(zip(cols, widths), start=1):
        cell = ws.cell(row=ws.max_row, column=i)
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = Alignment(vertical="center")
        ws.column_dimensions[cell.column_letter].width = w
    ws.freeze_panes = "A2"


def _summary_sheet(wb, project, bots):
    ws = wb.active
    ws.title = "Summary"
    _header(ws, ["Metric", "Value"], [28, 60])
    variables = sum(len(b["variables"]) for b in bots)
    packages = {p for b in bots for p in b["packages"]}
    vault_refs = sum(len(b["vault_references"]) for b in bots)
    rows = [
        ("Project", project["name"]),
        ("Source file", project["source_file"]),
        ("Task bots", project["task_bots"]),
        ("Active actions", project["active_actions"]),
        ("Disabled actions", project["disabled_actions"]),
        ("Variables", variables),
        ("Packages", len(packages)),
        ("Vault refs", vault_refs),
        ("Root bots", ", ".join(project["roots"]) or "—"),
    ]
    for r in rows:
        ws.append(r)
    for row in ws.iter_rows(min_row=2, min_col=1, max_col=1):
        row[0].font = LABEL_FONT


def _tree_rows(project, bots):
    """Flatten the nesting tree exactly as the UI walks it (roots, depth, orphans)."""
    rows, seen = [], set()

    def walk(name, depth):
        bot = next((x for x in bots if x["name"] == name), None)
        if bot is None:
            rows.append({"name": name, "depth": depth, "active": None, "disabled": None,
                         "complexity": None, "note": "not in export"})
            return
        if name in seen:
            return
        seen.add(name)
        called_by = bot["called_by"]
        note = "root" if bot["name"] in project["roots"] else ("orphan" if not called_by else "")
        rows.append({"name": name, "depth": depth, "active": bot["active_actions"],
                     "disabled": bot["disabled_actions"], "complexity": bot["complexity"], "note": note})
        for c in bot["calls"]:
            walk(c, depth + 1)

    for r in project["roots"]:
        walk(r, 0)
    for b in bots:
        if b["name"] not in seen and b["name"] not in project["roots"] and not any(
                b["name"] in y["calls"] for y in bots):
            rows.append({"name": b["name"], "depth": 0, "active": b["active_actions"],
                         "disabled": b["disabled_actions"], "complexity": b["complexity"], "note": "orphan"})
            seen.add(b["name"])
    return rows


def _tree_sheet(wb, project, bots):
    ws = wb.create_sheet("Call Tree")
    _header(ws, ["Task", "Depth", "Active", "Disabled", "Complexity"], [34, 8, 10, 12, 14])
    for r in _tree_rows(project, bots):
        ws.append([("  " * r["depth"]) + r["name"], r["depth"], r["active"], r["disabled"],
                   r["complexity"]])


def _bots_sheet(wb, bots):
    ws = wb.create_sheet("Task bots")
    _header(ws, ["Task bot", "Active Actions", "Disabled Actions", "Variables", "Top packages", "Score", "Complexity"],
            [34, 10, 12, 12, 46, 10, 14])
    for b in bots:
        ws.append([b["name"], b["active_actions"], b["disabled_actions"],
                   len(b["variables"]), ", ".join(list(b["packages"])[:4]),
                   b["complexity_score"], b["complexity"]])


def _detail_sheet(wb, bots):
    ws = wb.create_sheet("Task bot metrics")
    _header(ws, ["Task bot", "metrics", "Value"], [30, 24, 90])
    for b in bots:
        first = True
        fields = [
            ("Active / disabled", f"{b['active_actions']} / {b['disabled_actions']}"),
            ("Max nesting depth", b["max_nesting_depth"]),
            ("Complexity score", f"{b['complexity_score']} → {b['complexity']}"),
            ("Packages", ", ".join(f"{p} ({n})" for p, n in b["packages"].items()) or "—"),
            ("Variables", ", ".join(
                v["name"] + (" ⇦" if v["input"] else "") + (" ⇨" if v["output"] else "")
                for v in b["variables"]) or "—"),
            ("Calls", ", ".join(b["calls"]) or "—"),
            ("Called by", ", ".join(b["called_by"]) or "— (root)"),
            ("Vault refs", ", ".join(b["vault_references"]) or "—"),
            ("Global values", ", ".join(b["global_value_references"]) or "—"),
        ]
        for label, value in fields:
            ws.append([b["name"] if first else "", label, value])
            if first:
                ws.cell(row=ws.max_row, column=1).font = SECTION_FONT
                first = False
        ws.append(["", "", ""])


def build_workbook(project, bots):
    """Returns a seeked BytesIO of the .xlsx."""
    wb = Workbook()
    _summary_sheet(wb, project, bots)
    _tree_sheet(wb, project, bots)
    _bots_sheet(wb, bots)
    _detail_sheet(wb, bots)
    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return buf


# ---------------------------------------------------------------------------
# Interactive entry point
# ---------------------------------------------------------------------------
def _safe(name):
    return re.sub(r"[^A-Za-z0-9_-]+", "_", name) or "project"


def _default_input():
    default = os.path.join(os.path.dirname(os.path.abspath(__file__)),
                           "samples", "AR_CashApplication_Main.zip")
    if os.path.exists(default):
        return default
    return os.path.join(os.path.expanduser("~"), "Downloads",
                        "AR_CashApplication_Main.zip")


def main():
    print("=" * 62)
    print("  A360 -> Excel standalone converter")
    print("  Asks for an export, parses it, saves a 4-sheet .xlsx to your laptop.")
    print("=" * 62)
    print("Sheets: Summary | Call Tree | Task bots | Task bot metrics")
    print("Accepts: a Control Room export (.zip), a folder of bots, or one .json")
    print()

    default = _default_input()
    hint = f" (Enter = use sample: {default})" if os.path.exists(default) else ""
    raw = input(f"Path to .zip / folder / .json{hint}: ").strip().strip('"')
    path = raw or (default if os.path.exists(default) else "")
    if not path or not os.path.exists(path):
        print(f"ERROR: path not found — {path!r}"); sys.exit(1)

    print("Parsing ...")
    parsed = parse_export(path)
    bots = parsed["bots"]
    if not bots:
        print("ERROR: No task-bot JSON (objects with a 'nodes' array) found in the upload.")
        sys.exit(1)

    project = {
        "name": os.path.splitext(os.path.basename(path.rstrip("/\\")))[0],
        "source_file": os.path.basename(path),
        "task_bots": parsed["task_bots"],
        "active_actions": parsed["active_actions"],
        "disabled_actions": parsed["disabled_actions"],
        "roots": parsed["roots"],
    }

    default_out = os.path.join(os.path.expanduser("~"), "Downloads",
                               f"{_safe(project['name'])}_parsed.xlsx")
    raw_out = input(f"Output .xlsx path (Enter = Downloads folder): ").strip().strip('"')
    out = raw_out or default_out
    if not out.lower().endswith(".xlsx"):
        out += ".xlsx"
    out_dir = os.path.dirname(os.path.abspath(out))
    os.makedirs(out_dir, exist_ok=True)
    if os.path.exists(out):   # never silently overwrite — timestamp it
        stem, ext = os.path.splitext(out)
        out = f"{stem}_{datetime.now():%Y%m%d_%H%M%S}{ext}"

    print("Building workbook ...")
    with open(out, "wb") as fh:
        fh.write(build_workbook(project, bots).getvalue())

    print()
    print(f"  Task bots       : {project['task_bots']}")
    print(f"  Active actions  : {project['active_actions']}")
    print(f"  Disabled actions: {project['disabled_actions']}")
    print(f"  Root bots       : {', '.join(project['roots']) or '—'}")
    print()
    print(f"Saved -> {os.path.abspath(out)}")
    print("Done.")


if __name__ == "__main__":
    try:
        main()
    except KeyboardInterrupt:
        print("\nCancelled.")
        sys.exit(130)